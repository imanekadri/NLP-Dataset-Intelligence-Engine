"""Utilitaires pour l'Agent 1"""

import hashlib
import chardet
import re
from langdetect import detect, DetectorFactory
from langdetect.lang_detect_exception import LangDetectException
import pandas as pd
import json
from typing import Optional, Dict, Any, List
import os

# Pour des résultats reproductibles de détection de langue
DetectorFactory.seed = 42

class FileUtils:
    """Utilitaires pour la manipulation des fichiers"""
    
    @staticmethod
    def detect_file_type(file_path: str) -> str:
        """Détecte le type MIME du fichier sans magic"""
        ext = os.path.splitext(file_path)[1].lower()
        types = {
            '.csv': 'text/csv',
            '.txt': 'text/plain',
            '.json': 'application/json',
            '.pdf': 'application/pdf',
            '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            '.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            '.log': 'text/plain'
        }
        return types.get(ext, 'application/octet-stream')
    
    @staticmethod
    def detect_encoding(file_path: str) -> str:
        """Détecte l'encodage d'un fichier texte"""
        with open(file_path, 'rb') as f:
            raw_data = f.read()
        result = chardet.detect(raw_data)
        return result['encoding'] or 'utf-8'
    
    @staticmethod
    def compute_file_hash(file_path: str) -> str:
        """Calcule le hash MD5 du fichier"""
        hash_md5 = hashlib.md5()
        with open(file_path, "rb") as f:
            for chunk in iter(lambda: f.read(4096), b""):
                hash_md5.update(chunk)
        return hash_md5.hexdigest()
    
    @staticmethod
    def detect_language(text: str) -> Optional[str]:
        """Détecte la langue d'un texte"""
        try:
            return detect(text[:1000])
        except LangDetectException:
            return None

class TextExtractor:
    """Extraction de texte depuis différents formats"""
    
    @staticmethod
    def extract_from_csv(file_path: str, **kwargs) -> List[Dict]:
        """Extrait le texte d'un fichier CSV"""
        df = pd.read_csv(file_path, **kwargs)
        texts = []
        for _, row in df.iterrows():
            text = ' '.join([str(val) for val in row.values if pd.notna(val)])
            texts.append({
                'text': text,
                'metadata': {
                    'source': file_path,
                    'row': _,
                    'columns': list(row.index)
                }
            })
        return texts
    
    @staticmethod
    def extract_from_json(file_path: str) -> List[Dict]:
        """Extrait le texte d'un fichier JSON"""
        with open(file_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
        
        texts = []
        if isinstance(data, list):
            for i, item in enumerate(data):
                texts.append({
                    'text': json.dumps(item, ensure_ascii=False),
                    'metadata': {
                        'source': file_path,
                        'index': i,
                        'type': 'json_item'
                    }
                })
        else:
            texts.append({
                'text': json.dumps(data, ensure_ascii=False),
                'metadata': {
                    'source': file_path,
                    'type': 'json_object'
                }
            })
        return texts
    
    @staticmethod
    def extract_from_txt(file_path: str, encoding: str = 'utf-8') -> List[Dict]:
        """Extrait le texte d'un fichier texte"""
        with open(file_path, 'r', encoding=encoding) as f:
            content = f.read()
        
        paragraphs = [p.strip() for p in content.split('\n\n') if p.strip()]
        
        texts = []
        for i, para in enumerate(paragraphs):
            texts.append({
                'text': para,
                'metadata': {
                    'source': file_path,
                    'paragraph': i,
                    'total_paragraphs': len(paragraphs)
                }
            })
        return texts
    
    @staticmethod
    def extract_from_pdf(file_path: str) -> List[Dict]:
        """Extrait le texte d'un fichier PDF"""
        try:
            import PyPDF2
            texts = []
            with open(file_path, 'rb') as f:
                pdf_reader = PyPDF2.PdfReader(f)
                for page_num in range(len(pdf_reader.pages)):
                    page = pdf_reader.pages[page_num]
                    text = page.extract_text()
                    if text.strip():
                        texts.append({
                            'text': text,
                            'metadata': {
                                'source': file_path,
                                'page': page_num + 1,
                                'total_pages': len(pdf_reader.pages)
                            }
                        })
            return texts
        except ImportError:
            print("PyPDF2 non installé.")
            return []
    
    @staticmethod
    def extract_from_docx(file_path: str) -> List[Dict]:
        """Extrait le texte d'un fichier DOCX"""
        try:
            from docx import Document
            texts = []
            doc = Document(file_path)
            
            for i, para in enumerate(doc.paragraphs):
                if para.text.strip():
                    texts.append({
                        'text': para.text,
                        'metadata': {
                            'source': file_path,
                            'paragraph': i,
                            'type': 'paragraph'
                        }
                    })
            
            for table in doc.tables:
                for row in table.rows:
                    row_text = ' '.join([cell.text for cell in row.cells])
                    if row_text.strip():
                        texts.append({
                            'text': row_text,
                            'metadata': {
                                'source': file_path,
                                'type': 'table_row'
                            }
                        })
            return texts
        except ImportError:
            print("python-docx non installé.")
            return []
    
    @staticmethod
    def extract_from_log(file_path: str, encoding: str = 'utf-8') -> List[Dict]:
        """Extrait le texte d'un fichier log"""
        with open(file_path, 'r', encoding=encoding) as f:
            lines = f.readlines()
        
        texts = []
        for i, line in enumerate(lines):
            if line.strip():
                texts.append({
                    'text': line.strip(),
                    'metadata': {
                        'source': file_path,
                        'line': i + 1
                    }
                })
        return texts

class TextCleaner:
    """Nettoyage avancé du texte pour éliminer le bruit et les artefacts"""
    
    @staticmethod
    def clean(text: str) -> str:
        if not text:
            return ""
        
        # 1. Supprimer les artefacts de Jupyter Notebook (In [1]:, Out [1]:)
        text = re.sub(r'In\s*\[\d+\]:\s*', '', text)
        text = re.sub(r'Out\s*\[\d+\]:\s*', '', text)
        
        # 2. Supprimer les URLs
        text = re.sub(r'https?://\S+|www\.\S+', ' ', text)
        
        # 3. Supprimer les chemins de fichiers (Windows/Linux)
        text = re.sub(r'[a-zA-Z]:\\[\\\w\s.-]+', ' ', text)
        text = re.sub(r'/\w+/\w+/\S+', ' ', text)
        
        # 4. Supprimer les balises HTML résiduelles
        text = re.sub(r'<.*?>', '', text)
        
        # 5. Normaliser les espaces et sauts de ligne
        text = re.sub(r'\s+', ' ', text).strip()
        
        return text